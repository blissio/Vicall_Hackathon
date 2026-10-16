from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
root=Path('output/presentation'); source=(root/'Judge-QA-and-Technical-Study-Guide.md').read_text(encoding='utf-8')
doc=Document(); sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
for name,size in [('Normal',10.5),('Title',25),('Heading 1',16),('Heading 2',12),('Heading 3',11),('Subtitle',12)]:
 st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(6)
 if name.startswith('Heading'):st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run('Pittsburgh IQ  |  ');r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
def inline(p,line):
 for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)',line):
  if not part:continue
  r=p.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
  if part.startswith('**'):r.bold=True
  if part.startswith('`'):r.font.name='Consolas';r.font.size=Pt(9)
lines=source.splitlines();i=0
while i<len(lines):
 line=lines[i].strip();i+=1
 if not line:continue
 if line.startswith('# '):
  doc.add_paragraph('Pittsburgh IQ Judge Preparation',style='Title')
  doc.add_paragraph('Questions and answers and technical study guide',style='Subtitle')
  doc.add_paragraph('Rayane El Yasti, Oleg Hartman, Alexander Kearney, Noah White')
  continue
 if line.startswith('Prepared from '):
  doc.add_paragraph('Use this guide to prepare concise judge answers and understand how the application works. Practice the speaking answers first, then review the technical detail. Distinguish implemented features from measured results and proposed validation.');continue
 if line.startswith('## '):
  title=line[3:].replace('&','and');doc.add_paragraph(title,style='Heading 1');continue
 if line.startswith('### '):
  p=doc.add_paragraph(line[4:],style='Heading 2');continue
 if line.startswith('|'):
  rows=[line]
  while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
  rows=[x for x in rows if not re.fullmatch(r'[|\s:\-]+',x)]
  vals=[[x.strip() for x in row.strip('|').split('|')] for row in rows]
  table=doc.add_table(rows=0,cols=len(vals[0]));table.autofit=False
  widths=[2.3,4.7] if 'File' in vals[0] else [5.7,1.3]
  for rownum,row in enumerate(vals):
   cells=table.add_row().cells
   trpr=table.rows[-1]._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');trpr.append(el)
   if rownum==0:el=OxmlElement('w:tblHeader');trpr.append(el)
   for j,value in enumerate(row):
    c=cells[j];c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcpr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF0' if rownum==0 else 'FFFFFF');tcpr.append(sh)
    borders=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    tcpr.append(borders);mar=OxmlElement('w:tcMar')
    for edge in ['top','left','bottom','right']:
     el=OxmlElement('w:'+edge);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');mar.append(el)
    tcpr.append(mar);p=c.paragraphs[0];p.paragraph_format.space_after=Pt(3);inline(p,value)
    if rownum==0:
     for r in p.runs:r.bold=True
    if widths[j]<2:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  doc.add_paragraph().paragraph_format.space_after=Pt(2);continue
 p=doc.add_paragraph();inline(p,line)
doc.core_properties.title='Pittsburgh IQ Judge Preparation';doc.core_properties.subject='Judge questions and technical study guide';doc.core_properties.author='Pittsburgh IQ Team'
doc.save(root/'Pittsburgh-IQ-Judge-QA-Study-Guide.docx')
print('Saved document')
